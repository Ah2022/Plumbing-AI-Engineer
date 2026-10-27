"""Portable .plmcalc packages and native form definitions; no Excel runtime."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED,BadZipFile
from xml.etree.ElementTree import ParseError
import hashlib,json,math,re

from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries,get_column_letter
from openpyxl.formula.tokenizer import Tokenizer
from workbooks.reader import read_workbook
from workbooks.library import describe,run_scenario
from workbooks.formulas import Evaluator,FormulaError

FORMAT='plumbing-interactive-calculation'
VERSION=1
FORM_COMPILER_VERSION=2
LOCAL=Path(__file__).resolve().parents[1]/'local_sources'/'interactive_packages'
MAX_PACKAGE=15_000_000


def digest(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,separators=(',',':'),ensure_ascii=False,allow_nan=False).encode()).hexdigest()


def auxiliary(name):return name.lower().startswith(('data','reference','000','instruction'))


def row_index(cells):
    rows={}
    for key,c in cells.items():
        row,col=coordinate_to_tuple(key);rows.setdefault(row,[]).append((col,key,c))
    return rows


def label_for(cells,addr,rows=None,referenced=None):
    row,col=coordinate_to_tuple(addr)
    rows=rows or row_index(cells)
    labels=[]
    for k,key,c in rows.get(row,[]):
        v=c['value']
        if k<col and isinstance(v,str):
            if not v.strip() and c['formula']:
                v=next((s for s in re.findall(r'"([^"]+)"',c['formula']) if len(s)>3),'')
            if v.strip():labels.append((k,v))
    label=re.sub(r'\s+',' ',max(labels)[1]).strip() if labels else None
    header=None
    for r in sorted((r for r in rows if 9<=r<row),reverse=True):
        if any(key in (referenced or set()) for _,key,_ in rows[r]):continue
        items=[(k,c) for k,_,c in rows[r] if c['value'] is not None and str(c['value']).strip()]
        if len(items)>1 and all(isinstance(c['value'],str) and not c['formula'] for _,c in items):
            header=next((c['value'] for k,c in items if k==col),None)
            if header:break
    if label:return label+(' — '+re.sub(r'\s+',' ',header).strip() if header and header.strip()!=label else '')
    # Column headings in repeated schedule rows.
    for r in range(row-1,max(row-8,8),-1):
        c=cells.get(f'{get_column_letter(col)}{r}')
        if c and isinstance(c['value'],str) and c['value'].strip() and c['formula'] is None:
            return re.sub(r'\s+',' ',c['value']).strip()+f' (item {row})'
    return None


def _sections(cells):
    sections=[]
    rows={}
    for addr,c in cells.items():
        row,col=coordinate_to_tuple(addr);rows.setdefault(row,[]).append((col,c))
    for row,items in rows.items():
        if row<9:continue
        first=min(items,key=lambda x:x[0])
        nonempty=[c for _,c in items if c['value'] is not None and str(c['value']).strip()]
        if len(nonempty)==1 and isinstance(first[1]['value'],str) and not first[1]['formula'] and all(c['formula'] is None and isinstance(c['value'],str) for c in nonempty):
            text=re.sub(r'\s+',' ',first[1]['value']).strip()
            if len(text)<90 and text.lower()!='water tank sizing':sections.append((row,text))
    return sorted(sections)


def compile_forms(book):
    forms={}
    for name,sheet in book['sheets'].items():
        if auxiliary(name) or sheet['state']!='visible':continue
        cells=sheet['cells'];rows=row_index(cells);referenced=set();formulas=[(a,c) for a,c in cells.items() if c['formula']]
        for _,c in formulas:
            try:tokens=Tokenizer('='+c['formula']).items
            except Exception:continue
            for t in tokens:
                if t.type!='OPERAND' or t.subtype!='RANGE':continue
                text=t.value.replace('$','')
                if '!' in text:
                    sh,text=text.rsplit('!',1)
                    if sh.strip("'")!=name:continue
                if re.fullmatch('[A-Za-z]{1,3}[1-9][0-9]*',text):referenced.add(text.upper())
                elif re.fullmatch('[A-Za-z]{1,3}[1-9][0-9]*:[A-Za-z]{1,3}[1-9][0-9]*',text):
                    c1,r1,c2,r2=range_boundaries(text)
                    if (c2-c1+1)*(r2-r1+1)<=200:
                        referenced.update(f'{get_column_letter(c)}{r}' for r in range(r1,r2+1) for c in range(c1,c2+1))
        # Missing scalar inputs are retained as blank fields if they have a real label.
        for addr in sorted(referenced-set(cells),key=coordinate_to_tuple):
            row,col=coordinate_to_tuple(addr)
            is_merge_child=False
            for area in sheet.get('merges',[]):
                c1,r1,c2,r2=range_boundaries(area)
                if c1<=col<=c2 and r1<=row<=r2 and (col,row)!=(c1,r1):is_merge_child=True;break
            if 9<=row<=1000 and col<=100 and not is_merge_child and label_for(cells,addr,rows,referenced):
                cells[addr]=dict(value=None,formula=None,error=False,editable_hint=True)
        # Rebuild the row index after blank inputs are materialized. Otherwise the
        # first and second compilation can infer different table headings.
        rows=row_index(cells)
        sections=_sections(cells);fields=[];outputs=[];seen={}
        for addr,c in sorted(cells.items(),key=lambda item:coordinate_to_tuple(item[0])):
            row,col=coordinate_to_tuple(addr)
            if row<9 or col==1:continue
            label=label_for(cells,addr,rows,referenced)
            if not label:continue
            dynamic_labels=[(k,key) for k,key,item in rows.get(row,[]) if k<col and item['formula'] and isinstance(item['value'],str)]
            label_cell=max(dynamic_labels)[1] if dynamic_labels else None
            section=next((text for r,text in reversed(sections) if r<row),'Design inputs')
            if c['formula']:
                # Suppress decorative duplicates and formula-driven text headings.
                if isinstance(c['value'],str) and not c['error'] and not re.search(r'\d',c['value']):continue
                outputs.append(dict(cell=addr,label=label,section=section,label_cell=label_cell));continue
            if addr not in referenced and not c.get('validation'):continue
            if isinstance(c['value'],str) and not (c.get('validation') or c.get('editable_hint')):continue
            seen[label]=seen.get(label,0)+1
            if seen[label]>1:label+=f' ({seen[label]})'
            fields.append(dict(cell=addr,label=label,section=section,kind='boolean' if isinstance(c['value'],bool) else 'number' if c['value'] is None or isinstance(c['value'],(int,float)) else 'text',
                               percent='%' in c.get('number_format',''),validation=c.get('validation'),label_cell=label_cell))
        if len(fields)>500 or len(outputs)>1000:raise ValueError('Calculation form exceeds supported size')
        if not outputs:
            outputs=[dict(cell=a,label=f'Result {i+1}',section='Results') for i,(a,_) in enumerate(formulas[:100])]
        if outputs:forms[name]=dict(inputs=fields,outputs=outputs)
    return forms


def build_package(book,filename,metadata=None):
    book=deepcopy(book);book.pop('package_sha256',None)
    payload=dict(format=FORMAT,version=VERSION,form_compiler_version=FORM_COMPILER_VERSION,title=Path(filename).stem,
                 metadata=metadata or describe(book,filename),workbook=book)
    payload['forms']=compile_forms(book)
    payload['package_sha256']=digest(payload)
    return payload


def encode_package(package):return json.dumps(package,ensure_ascii=False,separators=(',',':'),allow_nan=False).encode('utf-8')


def validate_saved_controls(forms,book):
    """A legacy layout can change labels, never grant write access to formulas."""
    if not isinstance(forms,dict) or len(forms)>100:raise ValueError('Invalid saved form definitions')
    for sheet,form in forms.items():
        if sheet not in book['sheets'] or not isinstance(form,dict):raise ValueError('Invalid saved form worksheet')
        cells=book['sheets'][sheet]['cells']
        for group,limit in [('inputs',500),('outputs',1000)]:
            fields=form.get(group)
            if not isinstance(fields,list) or len(fields)>limit:raise ValueError('Invalid saved form definitions')
            for field in fields:
                if not isinstance(field,dict) or not isinstance(field.get('cell'),str):raise ValueError('Invalid saved form field')
                cell=cells.get(field['cell'])
                if cell is None or (group=='inputs' and cell.get('formula') is not None) or (group=='outputs' and cell.get('formula') is None):
                    raise ValueError('Saved form definitions do not match their source cells')


def decode_package(data):
    if len(data)>MAX_PACKAGE:raise ValueError('Interactive package exceeds 15 MB')
    def bad(value):raise ValueError('Non-finite value in interactive package')
    try:p=json.loads(data.decode('utf-8'),parse_constant=bad)
    except RecursionError as e:raise ValueError('Interactive package nesting is too deep') from e
    if not isinstance(p,dict) or p.get('format')!=FORMAT or p.get('version')!=VERSION:raise ValueError('Unsupported interactive calculation format/version')
    checksum=p.get('package_sha256');body={k:v for k,v in p.items() if k!='package_sha256'}
    if not isinstance(checksum,str) or checksum!=digest(body):raise ValueError('Interactive package checksum does not match its contents')
    compiler=p.get('form_compiler_version',1)
    if isinstance(compiler,bool) or not isinstance(compiler,int) or compiler<1 or compiler>FORM_COMPILER_VERSION:
        raise ValueError('Unsupported calculation layout version; update the engine to open this package')
    b=p.get('workbook',{})
    if not isinstance(b,dict):raise ValueError('Invalid calculation workbook')
    sheets=b.get('sheets')
    if not isinstance(p.get('title'),str) or not p['title'] or len(p['title'])>200:raise ValueError('Invalid calculation title')
    metadata=p.get('metadata')
    if not isinstance(metadata,dict):raise ValueError('Invalid calculation references')
    for key in ('primary_bases','all_bases','warnings','review_notes'):
        items=metadata.get(key,[])
        if not isinstance(items,list) or len(items)>1000 or any(not isinstance(v,str) or len(v)>2000 for v in items):raise ValueError('Invalid reference metadata')
    evidence=metadata.get('reference_evidence',[])
    if not isinstance(evidence,list) or len(evidence)>2000 or any(not isinstance(r,dict) or any(not isinstance(v,str) or len(v)>2000 for v in r.values()) for r in evidence):raise ValueError('Invalid reference evidence')
    if not isinstance(sheets,dict) or not 1<=len(sheets)<=100:raise ValueError('Invalid package worksheets')
    count=0
    for name,sh in sheets.items():
        if not isinstance(name,str) or not name or len(name)>100 or '!' in name:raise ValueError('Invalid worksheet name')
        if not isinstance(sh,dict) or not isinstance(sh.get('cells'),dict):raise ValueError('Invalid worksheet cells')
        merges=sh.get('merges',[])
        if not isinstance(merges,list) or len(merges)>5000:raise ValueError('Invalid merged-cell metadata')
        for area in merges:
            if not isinstance(area,str):raise ValueError('Invalid merged-cell reference')
            bounds=range_boundaries(area)
            if any(v is None for v in bounds):raise ValueError('Merged cells require bounded ranges')
        for addr,c in sh['cells'].items():
            count+=1
            if count>100000:raise ValueError('Package has too many cells')
            if not isinstance(addr,str) or not re.fullmatch('[A-Z]{1,3}[1-9][0-9]*',addr):raise ValueError('Invalid cell address')
            col,row,_,_=range_boundaries(addr)
            if col>16384 or row>1048576:raise ValueError('Cell exceeds workbook bounds')
            if not isinstance(c,dict):raise ValueError('Invalid cell record')
            value=c.get('value');formula=c.get('formula')
            if value is not None and not isinstance(value,(str,int,float,bool)):raise ValueError('Cell values must be scalar')
            if isinstance(value,str) and len(value)>16000:raise ValueError('Cell text too long')
            if formula is not None and (not isinstance(formula,str) or not formula or len(formula)>16000):raise ValueError('Invalid formula')
            if not isinstance(c.get('error',False),bool) or not isinstance(c.get('editable_hint',False),bool) or not isinstance(c.get('number_format',''),str):raise ValueError('Invalid field metadata')
            rule=c.get('validation')
            if rule is not None and (not isinstance(rule,dict) or rule.get('type')!='list' or not isinstance(rule.get('formula1'),str) or len(rule['formula1'])>16000):raise ValueError('Invalid input choices')
        if sh.get('state') not in ('visible','hidden','veryHidden'):raise ValueError('Invalid worksheet state')
    if not isinstance(b.get('names'),dict) or len(b['names'])>10000:raise ValueError('Invalid named references')
    if any(not isinstance(k,str) or not isinstance(v,str) or len(v)>16000 for k,v in b['names'].items()):raise ValueError('Invalid named reference')
    if not isinstance(b.get('sha256'),str) or not re.fullmatch('[0-9a-f]{64}',b['sha256']):raise ValueError('Missing original source identity')
    validate_saved_controls(p.get('forms'),b)
    # Validate without mutating the checksum-protected source record.
    forms=compile_forms(deepcopy(b))
    if compiler<FORM_COMPILER_VERSION:
        upgraded=build_package(b,p['title'],metadata)
        upgraded['title']=p['title']
        upgraded['upgraded_from_package']=checksum
        upgraded['package_sha256']=digest({k:v for k,v in upgraded.items() if k!='package_sha256'})
        return upgraded
    if p.get('forms')!=forms:raise ValueError('Interactive form definitions do not match their source cells')
    return p


def choices_for(book,sheet,field,overrides=None):
    rule=field.get('validation') or {};f=rule.get('formula1','').lstrip('=')
    if not f:return []
    try:
        if f.startswith('"') and f.endswith('"'):values=[v.strip() for v in f[1:-1].split(',')]
        else:
            ev=Evaluator(book,overrides,max_overrides=500);ref=ev.reference(f,sheet)
            values=[v for row in ev.matrix(ref) for v in row if v is not None]
        if len(values)>500:return []
        if field['kind']=='number':
            values=[float(v[:-1])/100 if isinstance(v,str) and v.endswith('%') else float(v) for v in values]
        return list(dict.fromkeys(values))
    except (ValueError,TypeError,FormulaError):return []


def import_bytes(data,filename):
    """Convert uploaded Excel or portable package/bundle without filesystem extraction."""
    if len(data)>110_000_000:raise ValueError('Upload exceeds 110 MB')
    if filename.lower().endswith('.plmcalc'):return [decode_package(data)]
    if filename.lower().endswith('.zip'):
        packages=[]
        with ZipFile(BytesIO(data)) as z:
            total=0
            for info in z.infolist():
                name=Path(info.filename.replace('\\','/')).name
                if name.startswith('~$') or not name.lower().endswith(('.xlsm','.xlsx','.plmcalc')):continue
                total+=info.file_size
                if info.file_size>30_000_000 or total>110_000_000:raise ValueError('Expanded library exceeds supported size')
                if len(packages)>=100:raise ValueError('Import at most 100 calculations per batch')
                packages.extend(import_bytes(z.read(info),name))
        if not packages:raise ValueError('No supported calculation files were found')
        return packages
    if not filename.lower().endswith(('.xlsx','.xlsm')):raise ValueError('Upload XLSX, XLSM, PLMCALC or ZIP')
    try:book=read_workbook(data)
    except (BadZipFile,ParseError,KeyError,IndexError) as e:raise ValueError('The file is not a readable XLSX/XLSM calculation workbook') from e
    return [build_package(book,filename)]


def bundle(packages):
    out=BytesIO()
    with ZipFile(out,'w',ZIP_DEFLATED) as z:
        for p in packages:
            slug=re.sub('[^A-Za-z0-9_-]+','-',p['title']).strip('-') or 'calculation'
            z.writestr(slug+'-'+p['package_sha256'][:12]+'.plmcalc',encode_package(p))
    return out.getvalue()


def install_project_packages(conn,pid,packages):
    if not conn.execute('SELECT 1 FROM projects WHERE project_id=?',(pid,)).fetchone():raise ValueError('Open a project first')
    validated=[decode_package(encode_package(p)) for p in packages]
    with conn:
        for p in validated:
            conn.execute('INSERT OR REPLACE INTO project_calculation_packages(project_id,package_sha256,title,package_json) VALUES (?,?,?,?)',
                         (pid,p['package_sha256'],p['title'],encode_package(p).decode('utf-8')))
    return len(validated)


def project_packages(conn,pid):
    return [decode_package(r[0].encode('utf-8')) for r in conn.execute('SELECT package_json FROM project_calculation_packages WHERE project_id=? ORDER BY title',(pid,))]


def run_package(p,sheet,values):
    form=p['forms'][sheet];allowed={f['cell'] for f in form['inputs']}
    if not isinstance(values,dict) or set(values)-allowed:raise ValueError('Only declared input fields can be changed')
    for f in form['inputs']:
        if f['cell'] not in values:continue
        value=values[f['cell']]
        if f['kind']=='number' and value is not None and (isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value)):
            raise ValueError(f['label']+' requires a finite number')
        if f['kind']=='boolean' and not isinstance(value,bool):raise ValueError(f['label']+' requires a yes/no value')
        if f['kind']=='text' and value is not None and not isinstance(value,str):raise ValueError(f['label']+' requires a text choice')
    book=deepcopy(p['workbook']);book['package_sha256']=p['package_sha256']
    outputs=form['outputs']
    # Reference methods with long schedules are evaluated in one shared, bounded evaluator.
    from workbooks.formulas import Evaluator,UnsupportedFormula
    overrides={sheet+'!'+cell:value for cell,value in values.items()}
    ev=Evaluator(book,overrides,max_overrides=500);results=[]
    for field in outputs:
        key=sheet+'!'+field['cell']
        try:
            value=ev.cell(sheet,field['cell']);status='CALCULATED_NOT_APPROVED';reason=''
        except UnsupportedFormula as e:value=None;status='NOT_EVALUATED';reason=str(e)
        except (FormulaError,RecursionError) as e:value=None;status='MISSING_DATA';reason=str(e)
        label=field['label']
        if field.get('label_cell'):
            try:label=ev.cell(sheet,field['label_cell']) or label
            except FormulaError:pass
        results.append(dict(cell=key,label=label,section=field['section'],value=value,status=status,reason=reason))
    request=dict(package_sha256=p['package_sha256'],worksheet=sheet,inputs=values)
    return dict(results=results,trace=ev.trace,inputs=values,source_sha256=book['sha256'],package_sha256=p['package_sha256'],
                input_sha256=digest(request),technical_status='CALCULATED_NOT_APPROVED' if all(r['status']=='CALCULATED_NOT_APPROVED' for r in results) else 'WARNING',
                review_status='REVIEW_REQUIRED',worksheet=sheet)
