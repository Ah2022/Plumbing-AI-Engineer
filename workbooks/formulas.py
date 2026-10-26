"""An explicit Excel subset for reference-sheet scenarios; never eval or VBA.

This reproduces formulas, including their source assumptions and mistakes.
It does not verify engineering criteria, macros, drawings or code compliance.
"""
from dataclasses import dataclass
from decimal import Decimal,ROUND_HALF_UP,ROUND_UP
import math
import operator
import re

from openpyxl.formula.tokenizer import Tokenizer,Token
from openpyxl.utils.cell import range_boundaries,get_column_letter,column_index_from_string


class FormulaError(ValueError): pass
class UnsupportedFormula(FormulaError): pass


@dataclass(frozen=True)
class Ref:
    sheet: str
    c1: int
    r1: int
    c2: int
    r2: int


class Parser:
    PRECEDENCE={':':90,'^':60,'*':50,'/':50,'+':40,'-':40,'&':30,'=':20,'<>':20,'<':20,'>':20,'<=':20,'>=':20}
    def __init__(self,formula):
        if len(formula)>16000:raise UnsupportedFormula('Formula too long')
        try:
            self.tokens=[]
            for t in Tokenizer('='+formula).items:
                if t.type=='WHITE-SPACE':continue
                # openpyxl lexes the dynamic range separator into ':INDIRECT('.
                if t.value.startswith(':') and t.type in ('FUNC','OPERAND'):
                    self.tokens.append(Token(':','OPERATOR-INFIX',''))
                    if len(t.value)>1:self.tokens.append(Token(t.value[1:],t.type,t.subtype))
                else:self.tokens.append(t)
        except Exception as e:raise UnsupportedFormula('Unparseable formula') from e
        self.i=0

    def parse(self):
        node=self.expression()
        if self.i!=len(self.tokens):raise UnsupportedFormula('Unsupported trailing formula syntax')
        return node

    def expression(self,minimum=0):
        if self.i>=len(self.tokens):raise FormulaError('Missing expression')
        t=self.tokens[self.i];self.i+=1
        if t.type=='OPERAND':node=('value',t.subtype,t.value)
        elif t.type=='OPERATOR-PREFIX':node=('unary',t.value,self.expression(70))
        elif t.type=='PAREN' and t.subtype=='OPEN':
            node=self.expression();self.close('PAREN')
        elif t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            while self.i<len(self.tokens) and self.tokens[self.i].subtype!='CLOSE':
                if self.tokens[self.i].type=='SEP':args.append(('literal',None))
                else:args.append(self.expression())
                if self.i<len(self.tokens) and self.tokens[self.i].type=='SEP':
                    self.i+=1
                    if self.i<len(self.tokens) and self.tokens[self.i].subtype=='CLOSE':args.append(('literal',None))
                else:break
            self.close('FUNC');node=('call',t.value[:-1].upper(),args)
        else:raise UnsupportedFormula('Unsupported token '+t.value)
        while self.i<len(self.tokens):
            t=self.tokens[self.i]
            if t.type=='OPERATOR-POSTFIX':
                if 80<minimum:break
                self.i+=1;node=('percent',node);continue
            if t.type!='OPERATOR-INFIX' or self.PRECEDENCE.get(t.value,-1)<minimum:break
            self.i+=1;node=('op',t.value,node,self.expression(self.PRECEDENCE[t.value]+1))
        return node

    def close(self,kind):
        if self.i>=len(self.tokens) or self.tokens[self.i].type!=kind or self.tokens[self.i].subtype!='CLOSE':
            raise FormulaError('Unclosed expression')
        self.i+=1


SUPPORTED={'IF','SUM','MAX','MIN','ISNUMBER','MATCH','INDEX','VLOOKUP','HLOOKUP','LOOKUP',
           'ROUNDUP','ROUND','EVEN','SQRT','ABS','RADIANS','SIN','COS','CHAR','COLUMN','INDIRECT'}


class Evaluator:
    def __init__(self,workbook,overrides=None):
        self.book=workbook;self.overrides={};self.cache={};self.failures={};self.active=set();self.parsed={};self.steps=0;self.trace={}
        if len(overrides or {})>100:raise ValueError('At most 100 input overrides')
        for key,value in (overrides or {}).items():
            sh,addr=self.address(key)
            cell=self.book['sheets'][sh]['cells'].get(addr)
            if cell is None or cell.get('formula') is not None:raise ValueError('Override must target an existing non-formula input: '+key)
            if not (value is None or isinstance(value,(str,bool,int,float))):raise ValueError('Input must be a scalar')
            if isinstance(value,(int,float)) and not math.isfinite(value):raise ValueError('Input must be finite')
            canonical=sh+'!'+addr
            if canonical in self.overrides:raise ValueError('Duplicate input cell alias: '+canonical)
            self.overrides[canonical]=value

    def address(self,key):
        if not isinstance(key,str) or '!' not in key:raise FormulaError('Use Sheet!A1 cell addresses')
        sh,addr=key.rsplit('!',1);sh=sh.strip("'").replace("''", "'");addr=addr.replace('$','').upper()
        if sh not in self.book['sheets'] or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*',addr):raise FormulaError('Unknown sheet/cell '+key)
        c1,r1,c2,r2=range_boundaries(addr)
        if c1>16384 or r1>1048576:raise FormulaError('Cell outside Excel bounds')
        return sh,addr

    def cell(self,sh,addr):
        key=sh+'!'+addr
        if key in self.overrides:return self.overrides[key]
        if key in self.cache:return self.cache[key]
        if key in self.failures:raise self.failures[key]
        if key in self.active:raise FormulaError('Circular reference at '+key)
        if len(self.active)>150:raise UnsupportedFormula('Dependency depth exceeds 150')
        self.steps+=1
        if self.steps>100000:raise UnsupportedFormula('Scenario dependency limit exceeded')
        c=self.book['sheets'][sh]['cells'].get(addr)
        if not c:self.cache[key]=None;return None
        if c.get('formula') is None:
            if c.get('error'):raise FormulaError(str(c['value'])+' at '+key)
            self.cache[key]=c['value'];return c['value']
        self.active.add(key)
        try:
            f=c['formula']
            if f not in self.parsed:self.parsed[f]=Parser(f).parse()
            value=self.scalar(self.eval(self.parsed[f],sh,addr))
            if value is None:value=0
            if isinstance(value,(int,float)) and not math.isfinite(value):raise FormulaError('Non-finite result')
            if isinstance(value,complex):raise FormulaError('Complex result unsupported')
            self.cache[key]=value;self.trace[key]=dict(formula=f,value=value)
            return value
        except FormulaError as e:self.failures[key]=e;raise
        except (ArithmeticError,ValueError,TypeError,IndexError) as e:
            error=FormulaError(f'Calculation error at {key}: {e}');self.failures[key]=error;raise error from e
        finally:self.active.remove(key)

    def reference(self,text,sh):
        if '[' in text or ']' in text:raise UnsupportedFormula('External workbook reference is unavailable')
        if '#REF!' in text:raise FormulaError('Broken #REF! reference')
        if '!' in text:
            sh,text=text.rsplit('!',1);sh=sh.strip("'").replace("''", "'")
        if sh not in self.book['sheets']:raise FormulaError('Unknown worksheet '+sh)
        text=text.replace('$','')
        if not re.fullmatch(r'[A-Za-z]{1,3}[1-9]\d*(?::[A-Za-z]{1,3}[1-9]\d*)?',text):
            name=text.upper();value=self.book['names'].get(sh+'!'+name,self.book['names'].get('!'+name))
            if value is None:raise UnsupportedFormula('Unknown name/range '+text)
            token=('name',sh,name)
            if token in self.active:raise FormulaError('Circular defined name '+name)
            self.active.add(token)
            try:return self.eval(Parser(value.lstrip('=')).parse(),sh,'A1')
            finally:self.active.remove(token)
        c1,r1,c2,r2=range_boundaries(text.upper())
        if c1>c2 or r1>r2 or c2>16384 or r2>1048576 or (c2-c1+1)*(r2-r1+1)>20000:
            raise UnsupportedFormula('Reference outside supported bounds')
        return Ref(sh,c1,r1,c2,r2)

    def scalar(self,value):
        if isinstance(value,Ref):
            if value.c1!=value.c2 or value.r1!=value.r2:raise UnsupportedFormula('Implicit intersection/array result unsupported')
            return self.cell(value.sheet,f'{get_column_letter(value.c1)}{value.r1}')
        return value

    def matrix(self,value):
        if isinstance(value,Ref):
            return [[self.cell(value.sheet,f'{get_column_letter(c)}{r}') for c in range(value.c1,value.c2+1)] for r in range(value.r1,value.r2+1)]
        return [[self.scalar(value)]]

    def number(self,value):
        value=self.scalar(value)
        if value is None:return 0.
        if isinstance(value,(int,float)):return float(value)
        try:return float(value)
        except (ValueError,TypeError):raise FormulaError('Expected number, found '+str(value))

    def truth(self,value):
        value=self.scalar(value)
        if value is None:return False
        if isinstance(value,(bool,int,float)):return bool(value)
        if isinstance(value,str) and value.upper() in ('TRUE','FALSE'):return value.upper()=='TRUE'
        raise FormulaError('Expected logical condition')

    @staticmethod
    def text(value):
        if value is None:return ''
        if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
        if isinstance(value,float) and value.is_integer():return str(int(value))
        return str(value)

    def eval(self,node,sh,addr):
        kind=node[0]
        if kind=='literal':return node[1]
        if kind=='value':
            typ,v=node[1:]
            if typ=='RANGE':return self.reference(v,sh)
            if typ=='NUMBER':return float(v)
            if typ=='TEXT':return v[1:-1].replace('""','"')
            if typ=='LOGICAL':return v=='TRUE'
            raise FormulaError(v)
        if kind=='unary':return self.number(self.eval(node[2],sh,addr))*(-1 if node[1]=='-' else 1)
        if kind=='percent':return self.number(self.eval(node[1],sh,addr))/100
        if kind=='op':
            op=node[1];a=self.eval(node[2],sh,addr);b=self.eval(node[3],sh,addr)
            if op==':':
                if not isinstance(a,Ref) or not isinstance(b,Ref) or a.sheet!=b.sheet:raise UnsupportedFormula('Invalid range operator')
                return self.reference(f'{get_column_letter(a.c1)}{a.r1}:{get_column_letter(b.c2)}{b.r2}',a.sheet)
            a=self.scalar(a);b=self.scalar(b)
            if op=='&':return self.text(a)+self.text(b)
            if op in ('=','<>','<','>','<=','>='):
                if a is None:a='' if isinstance(b,str) else 0
                if b is None:b='' if isinstance(a,str) else 0
                if isinstance(a,str):a=a.casefold()
                if isinstance(b,str):b=b.casefold()
                if type(a)!=type(b) and isinstance(a,str)!=isinstance(b,str):
                    if op in ('=','<>'):return op=='<>'
                    raise UnsupportedFormula('Mixed text/number ordering unsupported')
                return {'=':operator.eq,'<>':operator.ne,'<':operator.lt,'>':operator.gt,'<=':operator.le,'>=':operator.ge}[op](a,b)
            return {'+':operator.add,'-':operator.sub,'*':operator.mul,'/':operator.truediv,'^':operator.pow}[op](self.number(a),self.number(b))
        if kind=='call':return self.call(node[1],node[2],sh,addr)
        raise UnsupportedFormula('Unknown expression')

    def match(self,target,values,mode):
        # Numeric/text types remain distinct. Approximate lookup requires sorted data.
        target=self.scalar(target)
        if target is None:target=0
        def compatible(v):return isinstance(v,str)==isinstance(target,str) and v is not None and not isinstance(v,bool)
        def key(v):return v.casefold() if isinstance(v,str) else v
        pairs=[(i,v) for i,v in enumerate(values) if compatible(v)]
        if mode==0:
            if isinstance(target,str) and any(c in target for c in '*?~'):raise UnsupportedFormula('Wildcard MATCH requires an explicit adapter')
            for i,v in pairs:
                if key(v)==key(target):return i+1
        elif mode in (-1,1):
            vals=[key(v) for _,v in pairs]
            if vals!=sorted(vals,reverse=mode==-1):raise FormulaError('Approximate lookup vector is not sorted')
            eligible=[i+1 for i,v in pairs if (key(v)<=key(target) if mode==1 else key(v)>=key(target))]
            if eligible:return eligible[-1]
        else:raise FormulaError('Unsupported MATCH mode')
        raise FormulaError('#N/A: lookup has no matching entry')

    def call(self,name,args,sh,addr):
        if name not in SUPPORTED:raise UnsupportedFormula('Unsupported function '+name+'; no cached fallback')
        ev=lambda n:self.eval(args[n],sh,addr)
        if name=='IF':
            if len(args) not in (2,3):raise FormulaError('IF requires 2 or 3 arguments')
            return ev(1) if self.truth(ev(0)) else ev(2) if len(args)>2 else False
        if name=='ISNUMBER':
            try:v=self.scalar(ev(0));return isinstance(v,(int,float)) and not isinstance(v,bool)
            except UnsupportedFormula:raise
            except FormulaError:return False
        if name=='COLUMN':
            v=ev(0) if args else self.reference(addr,sh)
            if not isinstance(v,Ref):raise FormulaError('COLUMN requires a reference')
            return v.c1
        if name=='INDIRECT':
            if len(args)>1 and not self.truth(ev(1)):raise UnsupportedFormula('R1C1 INDIRECT unsupported')
            return self.reference(self.text(self.scalar(ev(0))),sh)
        if name=='INDEX':
            v=ev(0);r=int(self.number(ev(1)));c=int(self.number(ev(2))) if len(args)>2 else 1
            if not isinstance(v,Ref):raise UnsupportedFormula('INDEX requires a range')
            if len(args)==2 and v.r1==v.r2:c,r=r,1
            if not (1<=r<=v.r2-v.r1+1 and 1<=c<=v.c2-v.c1+1):raise FormulaError('#REF!: INDEX out of bounds')
            return Ref(v.sheet,v.c1+c-1,v.r1+r-1,v.c1+c-1,v.r1+r-1)
        if name in ('MATCH','LOOKUP','VLOOKUP','HLOOKUP'):
            target=ev(0);table=self.matrix(ev(1))
            if name in ('MATCH','LOOKUP'):
                if len(table)>1 and len(table[0])>1:raise UnsupportedFormula('Lookup requires a vector')
                vector=[r[0] for r in table] if len(table)>1 else table[0]
                mode=int(self.number(ev(2))) if name=='MATCH' and len(args)>2 else 1
                index=self.match(target,vector,mode)-1
                if name=='MATCH':return index+1
                if len(args)!=3:raise UnsupportedFormula('LOOKUP requires explicit result vector')
                output=self.matrix(ev(2));out=[r[0] for r in output] if len(output)>1 else output[0]
                if len(out)!=len(vector):raise FormulaError('LOOKUP vectors differ in length')
                return out[index]
            col=int(self.number(ev(2)))
            exact=not self.truth(ev(3)) if len(args)>3 else False
            vector=[r[0] for r in table] if name=='VLOOKUP' else table[0]
            index=self.match(target,vector,0 if exact else 1)-1
            if not 1<=col<=(len(table[0]) if name=='VLOOKUP' else len(table)):raise FormulaError('#REF!: lookup column/row out of bounds')
            return table[index][col-1] if name=='VLOOKUP' else table[col-1][index]
        values=[ev(i) for i in range(len(args))]
        if name in ('SUM','MAX','MIN'):
            numbers=[]
            for v in values:
                for row in self.matrix(v):
                    numbers.extend(x for x in row if isinstance(x,(int,float)) and not isinstance(x,bool))
            if name=='SUM':return sum(numbers)
            return (max(numbers) if name=='MAX' else min(numbers)) if numbers else 0
        x=self.number(values[0])
        if name in ('ROUND','ROUNDUP'):
            digits=int(self.number(values[1]));quantum=Decimal(10)**-digits
            if abs(digits)>100:raise UnsupportedFormula('Rounding precision exceeds supported bounds')
            return float(Decimal(str(x)).quantize(quantum,rounding=ROUND_UP if name=='ROUNDUP' else ROUND_HALF_UP))
        if name=='EVEN':return math.copysign(2*math.ceil(abs(x)/2),x)
        if name=='CHAR':
            if not 1<=int(x)<=255:raise FormulaError('CHAR outside supported byte range')
            return bytes([int(x)]).decode('cp1252')
        return {'SQRT':math.sqrt,'ABS':abs,'RADIANS':math.radians,'SIN':math.sin,'COS':math.cos}[name](x)
